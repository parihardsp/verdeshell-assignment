"""Document text extraction service utilizing PyMuPDF4LLM, PyMuPDF, and python-docx.

Extracts page-by-page text preserving Markdown tables, headings, lists,
and paragraph boundaries for optimal RAG semantic chunking and retrieval
across PDF, DOCX, DOC, CSV, TXT, and MD formats.
"""

import csv
import io
import logging
import re
import shutil
import subprocess
from dataclasses import dataclass
from pathlib import Path

import pymupdf
import pymupdf4llm

from app.core.exceptions import ValidationError
from app.core.logging import LogPhase

logger = logging.getLogger(__name__)


@dataclass
class ExtractedPage:
    """Represents text extracted from an individual page or section."""

    page_number: int
    text: str


@dataclass
class ExtractionResult:
    """Unified container for extracted document text and structural metadata."""

    file_name: str
    file_type: str
    total_pages: int
    pages: list[ExtractedPage]
    full_text: str


class DocumentExtractor:
    """Extracts structured text content and page-level metadata from diverse file formats."""

    @classmethod
    def extract_from_path(cls, file_path: Path) -> ExtractionResult:
        """Extract text from a saved disk file path based on its extension or magic bytes."""
        suffix = file_path.suffix.lower()

        # Check magic bytes for robust format detection
        try:
            with open(file_path, "rb") as f:
                header = f.read(8)
            if header.startswith(b"%PDF-"):
                return cls.extract_pdf(file_path)
            elif header.startswith(b"PK\x03\x04") or suffix == ".docx":
                return cls.extract_docx(file_path)
            elif header.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1") or suffix == ".doc":
                return cls.extract_doc(file_path)
        except Exception as exc:
            logger.debug(
                "Magic byte probe failed | file=%s error=%s",
                file_path.name,
                type(exc).__name__,
                extra={"phase": LogPhase.EXTRACTION},
            )

        if suffix == ".pdf":
            return cls.extract_pdf(file_path)
        elif suffix == ".docx":
            return cls.extract_docx(file_path)
        elif suffix == ".doc":
            return cls.extract_doc(file_path)
        elif suffix == ".csv":
            return cls.extract_csv(file_path)
        else:
            return cls.extract_plain_text(file_path)

    @classmethod
    def extract_docx(cls, file_path: Path) -> ExtractionResult:
        """Extract page-preserving structured text, headings, and Markdown tables from a .docx file."""
        # 1. Primary: PyMuPDF4LLM structured Markdown extraction with page chunks
        try:
            md_pages = pymupdf4llm.to_markdown(str(file_path), page_chunks=True)
            pages: list[ExtractedPage] = []
            full_text_parts: list[str] = []

            for idx, page_item in enumerate(md_pages):
                page_text = page_item.get("text", "").strip()
                page_num = page_item.get("metadata", {}).get("page_number", idx + 1)
                if page_text:
                    pages.append(ExtractedPage(page_number=page_num, text=page_text))
                    full_text_parts.append(page_text)

            if pages:
                logger.info(
                    "Structured extraction completed | engine=pymupdf4llm format=docx file=%s pages=%d",
                    file_path.name,
                    len(pages),
                    extra={"phase": LogPhase.EXTRACTION},
                )
                return ExtractionResult(
                    file_name=file_path.name,
                    file_type="docx",
                    total_pages=len(pages),
                    pages=pages,
                    full_text="\n\n".join(full_text_parts),
                )
        except Exception as exc:
            logger.warning(
                "Primary parser failed, attempting fallback | file=%s fallback=python-docx error=%s",
                file_path.name,
                type(exc).__name__,
                extra={"phase": LogPhase.EXTRACTION},
            )

        # 2. Secondary: python-docx structured document model (headings, paragraphs, tables)
        try:
            import docx

            doc = docx.Document(str(file_path))
            parts: list[str] = []

            for child in doc.element.body:
                if child.tag.endswith("p"):
                    p = docx.text.paragraph.Paragraph(child, doc)
                    text = p.text.strip()
                    if not text:
                        continue
                    style_name = (p.style.name or "").lower() if p.style else ""
                    if "heading 1" in style_name:
                        parts.append(f"# {text}")
                    elif "heading 2" in style_name:
                        parts.append(f"## {text}")
                    elif "heading 3" in style_name:
                        parts.append(f"### {text}")
                    elif "heading 4" in style_name:
                        parts.append(f"#### {text}")
                    elif "list bullet" in style_name or "bullet" in style_name:
                        parts.append(f"- {text}")
                    elif "list number" in style_name:
                        parts.append(f"1. {text}")
                    else:
                        parts.append(text)
                elif child.tag.endswith("tbl"):
                    table = docx.table.Table(child, doc)
                    if not table.rows:
                        continue
                    header_cells = [
                        c.text.strip().replace("\n", " ") for c in table.rows[0].cells
                    ]
                    dedup_header: list[str] = []
                    for c in header_cells:
                        if not dedup_header or c != dedup_header[-1]:
                            dedup_header.append(c)
                    if not dedup_header:
                        continue

                    md_lines = [
                        "| " + " | ".join(dedup_header) + " |",
                        "| " + " | ".join("---" for _ in dedup_header) + " |",
                    ]
                    for row in table.rows[1:]:
                        row_cells = [
                            c.text.strip().replace("\n", " ") for c in row.cells
                        ]
                        dedup_row: list[str] = []
                        for c in row_cells:
                            if not dedup_row or c != dedup_row[-1]:
                                dedup_row.append(c)
                        padded = dedup_row + [""] * (len(dedup_header) - len(dedup_row))
                        md_lines.append(
                            "| " + " | ".join(padded[: len(dedup_header)]) + " |"
                        )
                    parts.append("\n".join(md_lines))

            full_text = "\n\n".join(parts).strip()
            if full_text:
                logger.info(
                    "Structured extraction completed | engine=python-docx format=docx file=%s chars=%d",
                    file_path.name,
                    len(full_text),
                    extra={"phase": LogPhase.EXTRACTION},
                )
                return ExtractionResult(
                    file_name=file_path.name,
                    file_type="docx",
                    total_pages=1,
                    pages=[ExtractedPage(page_number=1, text=full_text)],
                    full_text=full_text,
                )
        except Exception as exc:
            logger.warning(
                "Secondary parser failed, attempting fallback | file=%s fallback=pymupdf error=%s",
                file_path.name,
                type(exc).__name__,
                extra={"phase": LogPhase.EXTRACTION},
            )

        # 3. Tertiary: Native PyMuPDF block extractor
        try:
            pages = []
            full_text_parts = []
            with pymupdf.open(str(file_path)) as doc:
                for idx, page in enumerate(doc):
                    text = page.get_text().strip()
                    if text:
                        pages.append(ExtractedPage(page_number=idx + 1, text=text))
                        full_text_parts.append(text)

            if pages:
                return ExtractionResult(
                    file_name=file_path.name,
                    file_type="docx",
                    total_pages=len(pages),
                    pages=pages,
                    full_text="\n\n".join(full_text_parts),
                )
        except Exception as exc:
            logger.error(
                "Native PyMuPDF failed | file=%s error=%s",
                file_path.name,
                type(exc).__name__,
                extra={"phase": LogPhase.EXTRACTION},
            )

        raise ValidationError(
            f"Failed to extract readable text from DOCX file '{file_path.name}'. "
            f"Please verify the file is not corrupted or password-protected."
        )

    @classmethod
    def extract_doc(cls, file_path: Path) -> ExtractionResult:
        """Extract text from legacy Word .doc files."""
        # 1. Primary: PyMuPDF native document reader
        try:
            with pymupdf.open(str(file_path)) as doc:
                pages: list[ExtractedPage] = []
                full_text_parts: list[str] = []
                for idx, page in enumerate(doc):
                    text = page.get_text().strip()
                    if text:
                        pages.append(ExtractedPage(page_number=idx + 1, text=text))
                        full_text_parts.append(text)
                if pages:
                    logger.info(
                        "Structured extraction completed | engine=pymupdf format=doc file=%s pages=%d",
                        file_path.name,
                        len(pages),
                        extra={"phase": LogPhase.EXTRACTION},
                    )
                    return ExtractionResult(
                        file_name=file_path.name,
                        file_type="doc",
                        total_pages=len(pages),
                        pages=pages,
                        full_text="\n\n".join(full_text_parts),
                    )
        except Exception as exc:
            logger.debug(
                "Extraction engine failed | engine=pymupdf format=doc file=%s error=%s",
                file_path.name,
                type(exc).__name__,
                extra={"phase": LogPhase.EXTRACTION},
            )

        # 2. Secondary: macOS textutil CLI
        if shutil.which("textutil"):
            try:
                proc = subprocess.run(
                    ["textutil", "-convert", "txt", str(file_path), "-stdout"],
                    capture_output=True,
                    text=True,
                    timeout=15,
                    check=False,
                )
                if proc.returncode == 0 and proc.stdout.strip():
                    clean_text = proc.stdout.strip()
                    logger.info(
                        "Structured extraction completed | engine=textutil format=doc file=%s chars=%d",
                        file_path.name,
                        len(clean_text),
                        extra={"phase": LogPhase.EXTRACTION},
                    )
                    return ExtractionResult(
                        file_name=file_path.name,
                        file_type="doc",
                        total_pages=1,
                        pages=[ExtractedPage(page_number=1, text=clean_text)],
                        full_text=clean_text,
                    )
            except Exception as exc:
                logger.debug(
                    "Extraction engine failed | engine=textutil format=doc file=%s error=%s",
                    file_path.name,
                    type(exc).__name__,
                    extra={"phase": LogPhase.EXTRACTION},
                )

        # 3. Tertiary: antiword or catdoc CLI if present in environment
        for cmd in ["antiword", "catdoc"]:
            if shutil.which(cmd):
                try:
                    proc = subprocess.run(
                        [cmd, str(file_path)],
                        capture_output=True,
                        text=True,
                        timeout=15,
                        check=False,
                    )
                    if proc.returncode == 0 and proc.stdout.strip():
                        clean_text = proc.stdout.strip()
                        return ExtractionResult(
                            file_name=file_path.name,
                            file_type="doc",
                            total_pages=1,
                            pages=[ExtractedPage(page_number=1, text=clean_text)],
                            full_text=clean_text,
                        )
                except Exception:
                    pass

        # 4. Fallback: Parse ASCII / UTF-16 text runs from OLE2 binary stream
        try:
            raw_bytes = file_path.read_bytes()
            ascii_runs = re.findall(rb"[\x20-\x7E\r\n\t]{4,}", raw_bytes)
            utf16_runs = re.findall(rb"(?:[\x20-\x7E][\x00]){4,}", raw_bytes)

            extracted_strings: list[str] = []
            for r in ascii_runs:
                s = r.decode("latin-1", errors="ignore").strip()
                if s and not s.startswith("Root Entry") and not s.startswith("WordDocument") and len(s) > 3:
                    extracted_strings.append(s)
            for r in utf16_runs:
                s = r.decode("utf-16-le", errors="ignore").strip()
                if s and not s.startswith("Root Entry") and not s.startswith("WordDocument") and len(s) > 3:
                    extracted_strings.append(s)

            joined = "\n\n".join(extracted_strings).strip()
            if len(joined) > 20:
                logger.info(
                    "Structured extraction completed | engine=ole2_stream format=doc file=%s chars=%d",
                    file_path.name,
                    len(joined),
                    extra={"phase": LogPhase.EXTRACTION},
                )
                return ExtractionResult(
                    file_name=file_path.name,
                    file_type="doc",
                    total_pages=1,
                    pages=[ExtractedPage(page_number=1, text=joined)],
                    full_text=joined,
                )
        except Exception as exc:
            logger.debug(
                "Extraction engine failed | engine=ole2_stream format=doc file=%s error=%s",
                file_path.name,
                type(exc).__name__,
                extra={"phase": LogPhase.EXTRACTION},
            )

        raise ValidationError(
            f"Could not extract readable text from legacy Word document '{file_path.name}'. "
            f"Please save or convert the file to .docx or .pdf and re-upload."
        )

    @classmethod
    def extract_pdf(cls, file_path: Path) -> ExtractionResult:
        """Extract page-by-page text from PDF preserving tables, paragraphs, and headings."""
        # 1. Primary: PyMuPDF4LLM structured Markdown extraction
        try:
            md_pages = pymupdf4llm.to_markdown(str(file_path), page_chunks=True)
            pages: list[ExtractedPage] = []
            full_text_parts = []

            for idx, page_item in enumerate(md_pages):
                page_text = page_item.get("text", "").strip()
                page_num = page_item.get("metadata", {}).get("page_number", idx + 1)
                if page_text:
                    pages.append(ExtractedPage(page_number=page_num, text=page_text))
                    full_text_parts.append(page_text)

            if pages:
                logger.info(
                    "Structured extraction completed | engine=pymupdf4llm format=pdf file=%s pages=%d",
                    file_path.name,
                    len(pages),
                    extra={"phase": LogPhase.EXTRACTION},
                )
                return ExtractionResult(
                    file_name=file_path.name,
                    file_type="pdf",
                    total_pages=len(pages),
                    pages=pages,
                    full_text="\n\n".join(full_text_parts),
                )
        except Exception as exc:
            logger.warning(
                "Primary parser failed, attempting fallback | file=%s fallback=pymupdf_tables error=%s",
                file_path.name,
                type(exc).__name__,
                extra={"phase": LogPhase.EXTRACTION},
            )

        # 2. Fallback: Native PyMuPDF with table detection & coordinate-sorted blocks
        try:
            pages = []
            full_text_parts = []
            with pymupdf.open(str(file_path)) as doc:
                total_pages = len(doc)
                for page_index in range(total_pages):
                    page = doc[page_index]
                    page_num = page_index + 1
                    page_content_parts = []

                    # Check for tables on page
                    try:
                        tabs = page.find_tables()
                        if tabs and tabs.tables:
                            for tab in tabs.tables:
                                md_table = tab.to_markdown()
                                if md_table:
                                    page_content_parts.append(md_table)
                    except Exception:
                        pass

                    # Extract layout-ordered text blocks (sorted top-to-bottom, left-to-right)
                    blocks = page.get_text("blocks")
                    text_blocks = [b for b in blocks if b[4].strip() and b[6] == 0]
                    text_blocks.sort(key=lambda b: (round(b[1] / 10) * 10, b[0]))
                    sorted_text = "\n\n".join(b[4].strip() for b in text_blocks)

                    if sorted_text:
                        page_content_parts.append(sorted_text)

                    final_page_text = "\n\n".join(page_content_parts).strip()
                    if final_page_text:
                        pages.append(ExtractedPage(page_number=page_num, text=final_page_text))
                        full_text_parts.append(final_page_text)

            if not full_text_parts:
                raise ValidationError(
                    f"PDF '{file_path.name}' contains no readable text or is image-only."
                )

            return ExtractionResult(
                file_name=file_path.name,
                file_type="pdf",
                total_pages=total_pages,
                pages=pages,
                full_text="\n\n".join(full_text_parts),
            )
        except ValidationError:
            raise
        except Exception as exc:
            logger.error(
                "Native PyMuPDF failed | file=%s error=%s",
                file_path.name,
                type(exc).__name__,
                extra={"phase": LogPhase.EXTRACTION},
            )
            raise ValidationError(f"Failed to extract text from PDF '{file_path.name}': {exc!s}")

    @classmethod
    def extract_csv(cls, file_path: Path) -> ExtractionResult:
        """Extract CSV data and convert directly into a structured Markdown table."""
        raw_bytes = file_path.read_bytes()
        text = ""
        for enc in ["utf-8", "utf-8-sig", "latin-1"]:
            try:
                text = raw_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue

        if not text:
            raise ValidationError(f"Could not decode CSV '{file_path.name}'.")

        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        if not rows:
            raise ValidationError(f"CSV '{file_path.name}' is empty.")

        # Convert to Markdown Table
        header = rows[0]
        md_table_lines = [
            "| " + " | ".join(h.strip() for h in header) + " |",
            "| " + " | ".join("---" for _ in header) + " |",
        ]
        for row in rows[1:]:
            padded = row + [""] * (len(header) - len(row))
            md_table_lines.append(
                "| " + " | ".join(c.strip() for c in padded[: len(header)]) + " |"
            )

        formatted_table = "\n".join(md_table_lines)
        return ExtractionResult(
            file_name=file_path.name,
            file_type="csv",
            total_pages=1,
            pages=[ExtractedPage(page_number=1, text=formatted_table)],
            full_text=formatted_table,
        )

    @classmethod
    def extract_plain_text(cls, file_path: Path) -> ExtractionResult:
        """Extract text from plain text or markdown files (.txt, .md)."""
        raw_bytes = file_path.read_bytes()

        # Guard: Never treat binary archives (e.g. zip, docx, pdf, ole2) as plain text!
        if raw_bytes.startswith(b"PK\x03\x04"):
            logger.info(
                "MIME header detected, re-routing | file=%s target=extract_docx",
                file_path.name,
                extra={"phase": LogPhase.EXTRACTION},
            )
            return cls.extract_docx(file_path)
        if raw_bytes.startswith(b"%PDF-"):
            logger.info(
                "MIME header detected, re-routing | file=%s target=extract_pdf",
                file_path.name,
                extra={"phase": LogPhase.EXTRACTION},
            )
            return cls.extract_pdf(file_path)
        if raw_bytes.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            logger.info(
                "MIME header detected, re-routing | file=%s target=extract_doc",
                file_path.name,
                extra={"phase": LogPhase.EXTRACTION},
            )
            return cls.extract_doc(file_path)

        # Standard Git binary check: text files do not contain null bytes
        if b"\x00" in raw_bytes[:4096]:
            raise ValidationError(
                f"File '{file_path.name}' is an unreadable binary file (contains null bytes)."
            )

        # Standard text encodings
        text = ""

        # Test UTF-8 first (strict)
        for enc in ["utf-8", "utf-8-sig"]:
            try:
                text = raw_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue

        # If not UTF-8, test fallback encodings with binary safety check
        if not text:
            for enc in ["latin-1", "cp1252"]:
                try:
                    candidate = raw_bytes.decode(enc)
                    # Binary sanity check: if more than 15% non-printable characters, reject
                    non_printable = sum(
                        1 for ch in candidate if ord(ch) < 9 or (13 < ord(ch) < 32)
                    )
                    if len(candidate) > 0 and (non_printable / len(candidate)) > 0.15:
                        continue
                    text = candidate
                    break
                except Exception:
                    continue

        if not text or not text.strip():
            raise ValidationError(
                f"File '{file_path.name}' is empty or contains unreadable binary data."
            )

        clean_text = text.strip()
        suffix = file_path.suffix.lstrip(".").lower() or "txt"
        return ExtractionResult(
            file_name=file_path.name,
            file_type=suffix,
            total_pages=1,
            pages=[ExtractedPage(page_number=1, text=clean_text)],
            full_text=clean_text,
        )
